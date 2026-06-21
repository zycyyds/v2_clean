from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt

BASE = Path('/Users/wwz/Downloads/medical-agent-pipeline-main/v2')
OUT = BASE / 'v2_codebase_and_agent_framework.pptx'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(15, 23, 42)
TEAL = RGBColor(13, 148, 136)
CYAN = RGBColor(8, 145, 178)
SLATE = RGBColor(71, 85, 105)
LIGHT = RGBColor(248, 250, 252)
MID = RGBColor(226, 232, 240)
TEXT = RGBColor(30, 41, 59)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(22, 163, 74)
AMBER = RGBColor(217, 119, 6)
RED = RGBColor(220, 38, 38)


def add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None, dark=False):
    color = WHITE if dark else NAVY
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(0.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = color
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.72), Inches(1.05), Inches(11.8), Inches(0.45))
        tf2 = sub.text_frame
        tf2.clear()
        p2 = tf2.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(11)
        r2.font.color.rgb = WHITE if dark else SLATE


def add_footer(slide, text, dark=False):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(7.0), Inches(12), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(9)
    r.font.color.rgb = MID if dark else SLATE
    p.alignment = PP_ALIGN.RIGHT


def add_bullets(slide, x, y, w, h, items, font_size=16, color=TEXT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.clear()
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if isinstance(item, tuple):
            level, text = item
        else:
            level, text = 0, item
        p.text = text
        p.level = level
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(6)
        if level == 0:
            p.bullet = True


def add_card(slide, x, y, w, h, title, body_lines, accent=TEAL):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = MID
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.14), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    title_box = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.15), Inches(w - 0.35), Inches(0.35))
    p = title_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = NAVY
    add_bullets(slide, x + 0.28, y + 0.55, w - 0.4, h - 0.65, body_lines, font_size=11)


def add_metric(slide, x, y, w, h, number, label, color=TEAL):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = MID
    num_box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.12), Inches(w - 0.2), Inches(0.45))
    p = num_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = str(number)
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = color
    lab_box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.62), Inches(w - 0.2), Inches(0.35))
    p2 = lab_box.text_frame.paragraphs[0]
    r2 = p2.add_run()
    r2.text = label
    r2.font.size = Pt(10)
    r2.font.color.rgb = SLATE


def add_pipeline(slide, steps, y=2.1):
    x = 0.6
    box_w = 1.45
    gap = 0.15
    for i, (title, subtitle, color) in enumerate(steps, start=1):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(box_w), Inches(1.45))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = color
        circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.08), Inches(y + 0.12), Inches(0.34), Inches(0.34))
        circ.fill.solid()
        circ.fill.fore_color.rgb = color
        circ.line.fill.background()
        tb1 = slide.shapes.add_textbox(Inches(x + 0.48), Inches(y + 0.07), Inches(0.8), Inches(0.25))
        p1 = tb1.text_frame.paragraphs[0]
        r1 = p1.add_run()
        r1.text = str(i)
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = SLATE
        tb2 = slide.shapes.add_textbox(Inches(x + 0.12), Inches(y + 0.45), Inches(box_w - 0.2), Inches(0.28))
        p2 = tb2.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = title
        r2.font.size = Pt(12)
        r2.font.bold = True
        r2.font.color.rgb = NAVY
        tb3 = slide.shapes.add_textbox(Inches(x + 0.12), Inches(y + 0.74), Inches(box_w - 0.22), Inches(0.58))
        tf3 = tb3.text_frame
        tf3.word_wrap = True
        p3 = tf3.paragraphs[0]
        r3 = p3.add_run()
        r3.text = subtitle
        r3.font.size = Pt(9)
        r3.font.color.rgb = SLATE
        if i < len(steps):
            line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(x + box_w + 0.02), Inches(y + 0.52), Inches(0.11), Inches(0.28))
            line.fill.solid()
            line.fill.fore_color.rgb = MID
            line.line.fill.background()
        x += box_w + gap


# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.22))
accent.fill.solid(); accent.fill.fore_color.rgb = TEAL; accent.line.fill.background()
add_title(slide, 'Medical Agent Pipeline v2', '代码库核心技术、Agent 框架与结果生成路径解读', dark=True)
box = slide.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(7.2), Inches(2.4))
tf = box.text_frame
p = tf.paragraphs[0]
r = p.add_run(); r.text = '从“多 agent 状态机”重构为“单 PlannerReAct + Skill Registry + 分步代码执行”'; r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = WHITE
p2 = tf.add_paragraph(); p2.text = '目标：让 Agent 自主规划、多步试错、逐步沉淀中间产物，并把医疗多模态数据整理为可训练数据集。'; p2.font.size = Pt(15); p2.font.color.rgb = MID; p2.space_before = Pt(14)
add_metric(slide, 8.5, 1.8, 1.6, 1.0, '1', 'Planner Agent')
add_metric(slide, 10.2, 1.8, 1.4, 1.0, '37', 'README 声明技能数')
add_metric(slide, 11.7, 1.8, 1.0, 1.0, '2', '双阶段模式')
add_card(slide, 8.45, 3.1, 4.0, 2.4, '本次 PPT 聚焦', [
    'v2/ 的目录与职责划分',
    'Planner / DataExplorer / FeatureEngineer 的协作方式',
    'Skill 自动注册机制与运行时设计',
    'output/ 中结果如何一步步生成',
], accent=CYAN)
add_footer(slide, 'Source: v2/README.md, main.py, agent/, skills/, lib/, output/code_outputs', dark=True)

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, '1. 代码库总览：v2 想解决什么问题', '核心思想：保留 v1 step 能力，但把编排权交给更轻量的 ReAct + skill 层')
add_card(slide, 0.6, 1.45, 4.0, 2.2, 'v1 的问题', [
    'router / handoff / supervisor / state machine 逻辑较重',
    '每个 step 都绑定独立 agent，编排代码多、耦合高',
    '失败恢复依赖显式 ticket 和 resume 状态',
], accent=RED)
add_card(slide, 4.75, 1.45, 4.0, 2.2, 'v2 的改法', [
    '入口统一到 main.py',
    'skill 粒度下沉：把工具能力拆成原子技能',
    'PlannerReActAgent 通过 thought -> tool -> observe 自主推进',
], accent=TEAL)
add_card(slide, 8.9, 1.45, 3.8, 2.2, '直接收益', [
    '编排层更短、更易扩展',
    '新增 skill 不必改 planner 主逻辑',
    '失败后可在同一 ReAct 循环里直接修复并重试',
], accent=CYAN)
add_pipeline(slide, [
    ('入口', 'main.py 接受自然语言任务', TEAL),
    ('规划', 'LLM 决定先看什么、做什么', CYAN),
    ('执行', '调用 skill / runtime / Python 代码', TEAL),
    ('观察', '根据 SUCCESS / NEEDS_REPAIR 调整', AMBER),
    ('产出', '沉淀中间 CSV 与最终训练集', GREEN),
], y=4.2)
add_footer(slide, 'Source: v2/README.md:3-48, v2/main.py:1-18')

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, '2. 目录与分层职责', '代码不是按模型分，而是按“规划层 / skill 层 / runtime 层 / 输出层”拆分')
add_card(slide, 0.55, 1.4, 3.0, 4.9, 'agent/', [
    'planner_agent.py：构造单 PlannerReActAgent',
    'system_prompt.py：定义 DataExplorer / FeatureEngineer 规则',
    'two_phase_agents.py：双阶段工厂',
], accent=TEAL)
add_card(slide, 3.75, 1.4, 3.0, 4.9, 'skills/', [
    '_registry.py 自动发现 skill.py 并注册',
    '_common.py 注入路径与默认目录',
    '每个 skill 目录包含 SKILL.md + skill.py',
    '本仓实际可见 21 个 skill 目录，README 说明目标是 37 个原子 skill',
], accent=CYAN)
add_card(slide, 6.95, 1.4, 3.0, 4.9, 'lib/', [
    'agent_runtime.py：模型/消息/工具结果封装',
    'step23_runtime.py：目录扫描、OCR、文本抽取、宽表拼接',
    'step4~step7_runtime.py：选列、清洗、一致性校验、数据集导出',
], accent=AMBER)
add_card(slide, 10.15, 1.4, 2.6, 4.9, 'output/', [
    'logs/：完整运行日志',
    'code_outputs/：step1-step8 与最终 CSV',
    '这里是“结果如何得到”的最好证据链',
], accent=GREEN)
add_footer(slide, 'Source: directory scan, v2/skills/_registry.py, v2/lib/*.py')

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, '3. Agent 框架主线：单 Planner 与双阶段模式并存', '同一套 skill 底座，支持两种使用方式')
add_card(slide, 0.7, 1.5, 5.95, 4.7, '模式 A：单 PlannerReActAgent', [
    '入口：main.py -> run_one_shot()',
    'create_planner_agent() 创建一个 ReActAgent，并挂上全部已注册 skill',
    '适合：自由文本任务，一次性从探索走到结果',
    '行为：让 LLM 自己决定先扫描、再采样、再写 Python 处理代码',
], accent=TEAL)
add_card(slide, 6.8, 1.5, 5.8, 4.7, '模式 B：DataExplorer + FeatureEngineer', [
    '入口：main.py --two-phase -> run_two_phase()',
    '阶段1先全面理解数据，阶段2再分步做特征工程',
    '优点：更像“先做分析报告，再按计划施工”',
    '局限：当前实现里 toolkit 实际仍是 full toolkit，只是 prompt 中做了子集约束',
], accent=CYAN)
add_pipeline(slide, [
    ('Input', '自然语言任务 + 数据路径', TEAL),
    ('Agent Factory', '单 Planner 或双阶段工厂', CYAN),
    ('Toolkit', 'load_all_skills() 注册能力', TEAL),
    ('ReAct Loop', '多轮推理/调用/修复', AMBER),
    ('Artifacts', '日志 + CSV + JSON', GREEN),
], y=6.35)
add_footer(slide, 'Source: v2/main.py:102-155, v2/agent/planner_agent.py:56-82, v2/agent/two_phase_agents.py:81-112')

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, '4. Skill Registry：把工具能力变成可规划的技能清单', 'skill 不只是函数注册，还带有给 Agent 看的元数据清单')
add_card(slide, 0.7, 1.5, 4.0, 4.9, '自动发现机制', [
    '遍历 v2/skills/*/skill.py',
    '每个模块必须暴露 SKILL 字典 + register(toolkit)',
    '启动时统一 import 并注册到 Toolkit',
], accent=TEAL)
add_card(slide, 4.9, 1.5, 3.9, 4.9, 'Manifest 作用', [
    'render_skill_manifest() 按 explore / process / label / util 分层输出',
    'system prompt 里只放“一行摘要”',
    'Agent 无需先读整份文档，就知道每个 skill 何时使用',
], accent=CYAN)
add_card(slide, 9.0, 1.5, 3.6, 4.9, '代表性 skill', [
    'list_data_files / sample_table：探索输入数据结构',
    'run_python_code：执行自写 Python 代码',
    'select_features / export_ml_dataset：衔接 step4 / step7',
], accent=AMBER)
quote = slide.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.8), Inches(0.35))
p = quote.text_frame.paragraphs[0]
r = p.add_run(); r.text = '本质上：Planner 看到的不是“源代码文件”，而是“可以被规划和组合的技能目录”。'; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = NAVY
add_footer(slide, 'Source: v2/skills/_registry.py:1-85, sample skill files in v2/skills/*/skill.py')

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, '5. Runtime 设计：为什么 run_python_code 很关键', '这是 v2 最灵活的一层：让 Agent 用 Python 直接“落地数据处理”')
add_card(slide, 0.65, 1.45, 4.0, 4.9, 'run_python_code skill', [
    '执行任意 Python 字符串',
    '自动注入 OUTPUT_DIR',
    '收集 stdout / stderr / output_files',
    '返回 SUCCESS 或 NEEDS_REPAIR 供 ReAct 下一轮判断',
], accent=TEAL)
add_card(slide, 4.85, 1.45, 4.0, 4.9, 'lib/runtime 的价值', [
    '把 step2-3, step4, step5, step6, step7 的复杂逻辑封装成稳定函数',
    '既能被 skill 调用，也能单独运行',
    '大量地方显式输出 shape / 缺失 / 产物路径，方便 Agent 观察',
], accent=CYAN)
add_card(slide, 9.05, 1.45, 3.55, 4.9, '修复闭环', [
    '代码报错 -> skill 返回 NEEDS_REPAIR',
    'Agent 读取 traceback，改写代码后再试一次',
    '日志里能看到真实的“试错-修复-继续”过程',
], accent=AMBER)
add_footer(slide, 'Source: v2/skills/run_python_code/skill.py:20-89, v2/lib/step4_runtime.py, step5_runtime.py, step7_runtime.py')

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, '6. Step Runtime 示例：从结构化表到训练集', '并不是所有特征都靠 LLM 生造，很多是用稳定的数据处理模板完成的')
add_card(slide, 0.6, 1.55, 3.1, 4.7, 'Step2-3', [
    '可识别 directory / ocr_fill 两种模式',
    '扫描目录、OCR、文本抽取、再宽表拼接',
    '输出给 step4 的 next_input/input.csv',
], accent=TEAL)
add_card(slide, 3.95, 1.55, 2.8, 4.7, 'Step4', [
    '任务驱动列筛选',
    '生成 filtered.csv + selection_report.json',
    '并做输出一致性验证',
], accent=CYAN)
add_card(slide, 6.95, 1.55, 2.8, 4.7, 'Step5', [
    '列画像 + 风险分级 + 数据清洗',
    '对高风险列可尝试 LLM 生成清洗器',
    '保持行列结构不变',
], accent=AMBER)
add_card(slide, 9.95, 1.55, 2.8, 4.7, 'Step6/7', [
    'Step6 做患者一致性验证',
    'Step7 做标签构造、ICD 映射与导出',
    '最终给出训练集与模型配置',
], accent=GREEN)
add_footer(slide, 'Source: v2/lib/step23_runtime.py, step4_runtime.py, step5_runtime.py, step6_runtime.py, step7_runtime.py')

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, '7. 本项目实际产物：两条“结果生成路径”并行存在', 'output/ 里能看到同一任务被不同 Agent 使用方式完成过')
add_card(slide, 0.7, 1.55, 5.8, 4.8, '路径 A：单 Planner 直接生成简化训练集', [
    '日志：汇报.log',
    '先 list_data_files + sample_table 探索 9 张表',
    '随后一次 run_python_code 直接做聚合合并',
    '产物：hospital_los_dataset.csv（400 x 17）',
    '标签来自 admissions 的 los_hours，特征较粗粒度',
], accent=TEAL)
add_card(slide, 6.7, 1.55, 5.9, 4.8, '路径 B：双阶段模式逐步扩展特征宽表', [
    '日志：run_20260530_172242.log',
    'DataExplorer 先出 feature plan',
    'FeatureEngineer 按 step1-step9 逐个写中间 CSV',
    '产物：ml_dataset_final.csv（400 x 118）',
    '标签来自 icustays.los，多模态特征更丰富',
], accent=CYAN)
add_footer(slide, 'Source: output/logs/汇报.log, output/logs/run_20260530_172242.log, output/code_outputs/*.csv')

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, '8. DataExplorer：先理解数据，再给出特征施工图', '它做的不是训练，而是“把表结构、关联键、候选特征、标签来源讲清楚”')
add_metric(slide, 0.8, 1.45, 1.5, 0.95, '9', '扫描数据文件')
add_metric(slide, 2.45, 1.45, 2.0, 0.95, '1', '目标标签建议: icustays.los')
add_metric(slide, 4.65, 1.45, 2.0, 0.95, '3', '高相关核心表')
add_metric(slide, 6.85, 1.45, 2.0, 0.95, '6+', '模态来源')
add_card(slide, 0.7, 2.75, 4.0, 3.5, '它识别出的主表层次', [
    'icustays：核心标签 los',
    'admissions：住院上下文',
    'patients：人口统计学',
    'diagnoses_icd / labevents / prescriptions：结构化扩展特征',
], accent=TEAL)
add_card(slide, 4.9, 2.75, 3.7, 3.5, '它给出的处理策略', [
    '诊断：top-N ICD one-hot / 频次',
    '化验：按 itemid 聚合 mean/max/min',
    '用药：时长统计 + 类别计数',
    '文本/影像：先做轻量特征，后续可增强 NLP',
], accent=CYAN)
add_card(slide, 8.8, 2.75, 3.8, 3.5, '暴露出的真实问题', [
    '报告路径在一次运行中写入位置不一致',
    'FeatureEngineer 曾因找不到 report_path 而中断',
    '说明双阶段模式仍有工程打磨空间',
], accent=AMBER)
add_footer(slide, 'Source: run_20260530_171740.log, run_20260530_172242.log')

# Slide 10
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, '9. FeatureEngineer：结果是如何一步步得到的', '下面这条链路就是 output/code_outputs 里的真实证据')
steps = [
    ('step1', 'admissions 主表\n395 x 16', TEAL),
    ('step2', '+ icustays\n400 x 22', CYAN),
    ('step3', '+ patients\n400 x 25', TEAL),
    ('step4', '+ diagnosis\n400 x 45', AMBER),
    ('step5', '+ lab agg\n400 x 105', GREEN),
    ('step6', '+ prescriptions\n400 x 111', CYAN),
    ('step7', '+ CXR\n400 x 116', TEAL),
    ('step8', '+ notes\n400 x 118', AMBER),
]
add_pipeline(slide, steps, y=2.0)
add_card(slide, 0.8, 4.2, 5.7, 2.0, '关键扩展点', [
    'step4：诊断 top20 ICD one-hot，列数 25 -> 45',
    'step5：lab top20 itemid × 3 统计，列数 45 -> 105',
    'step8：文本长度类特征把宽表补到 118 列',
], accent=TEAL)
add_card(slide, 6.8, 4.2, 5.7, 2.0, '可见的工程特征', [
    '步骤 4 先后经历两次报错，再修复成功',
    '每步都打印 shape，便于 Agent 判断是否继续',
    '最终保留了完整的 step1-step8 中间 CSV',
], accent=CYAN)
add_footer(slide, 'Source: output/code_outputs summary, run_20260530_172242.log')

# Slide 11
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title(slide, '10. 最终结果对比：简化版 vs 多阶段宽表版', '两份结果都能训练，但表达能力和可追溯性明显不同')
add_card(slide, 0.8, 1.7, 5.5, 4.7, 'hospital_los_dataset.csv', [
    '维度：400 行 × 17 列',
    '标签：los_hours（由 admissions 时间差计算）',
    '特征：住院类型、人口学、诊断数量、lab_count、lab_value_mean、presc_count、icu_los',
    '优点：简单、快、容易解释',
    '局限：高维结构化信息被强烈压缩',
], accent=TEAL)
add_card(slide, 6.95, 1.7, 5.6, 4.7, 'ml_dataset_final.csv', [
    '维度：400 行 × 118 列',
    '标签：los（来自 icustays）',
    '特征：ICD one-hot + lab 多统计 + prescriptions 统计 + 影像类别 + 文本长度',
    '优点：信息密度高，更接近真实 ML 宽表',
    '代价：标签缺失用 0 填充、编码策略偏粗，需要再打磨',
], accent=CYAN)
add_footer(slide, 'Source: output/code_outputs/hospital_los_dataset.csv, ml_dataset_final.csv')

# Slide 12
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
add_title(slide, '11. 技术总结与下一步改进建议', 'v2 已经证明“Agent + Skill + Runtime”这条路线可行，但还可以更稳、更像生产系统', dark=True)
add_card(slide, 0.8, 1.65, 3.8, 4.6, '已经做对的事', [
    '把编排复杂度从多状态机收敛到单 ReAct 框架',
    'skill 注册机制让扩展新能力更自然',
    '中间 CSV + 日志构成了强可追踪的证据链',
], accent=TEAL)
add_card(slide, 4.8, 1.65, 3.8, 4.6, '当前暴露的风险', [
    'README 的 37 skills 与目录中实际 21 skills 存在差异',
    '双阶段 report_path 传递有路径不一致问题',
    '最终 los 缺失被填 0，可能污染监督信号',
], accent=AMBER)
add_card(slide, 8.8, 1.65, 3.8, 4.6, '建议优先级', [
    '统一 report / artifact 路径协议',
    '把 label 缺失处理改成显式过滤或单独标签策略',
    '进一步把文本和影像特征从“长度/标签”升级为真实表示学习',
], accent=CYAN)
add_footer(slide, 'Prepared from v2 source code and output artifacts', dark=True)

prs.save(str(OUT))
print(OUT)
