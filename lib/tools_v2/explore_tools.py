# -*- coding: utf-8 -*-
"""
探索类工具（Explore Tools）

供 ReActAgent 调用的数据探索工具，每个函数都是独立的原子操作。
Agent 通过多次调用这些工具来了解数据情况，再决定后续处理策略。

工具列表：
  - detect_input_type      检测输入数据类型
  - read_csv_sample        读取 CSV 样本及列分析
  - get_csv_full_info      获取 CSV 完整分析报告
  - scan_directory         扫描目录结构（患者分组）
  - list_directory         列出目录内容
  - read_text_sample       读取文本文件样本
"""
import os
import csv
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Set

# 导入本包配置
import sys
_HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

from config import (
    DataType, IMAGE_EXTENSIONS, TEXT_EXTENSIONS, CSV_EXTENSIONS, EXCEL_EXTENSIONS,
    JSONL_EXTENSIONS, MEDICAL_COL_KEYWORDS, get_skip_folders,
)


# ---------------------------------------------------------------------------
# 内部辅助
# ---------------------------------------------------------------------------

def _ext(path: str) -> str:
    return Path(path).suffix.lower()


def _is_image(path: str) -> bool:
    return _ext(path) in IMAGE_EXTENSIONS


def _is_csv(path: str) -> bool:
    return _ext(path) in CSV_EXTENSIONS


def _is_text(path: str) -> bool:
    return _ext(path) in TEXT_EXTENSIONS


def _is_excel(path: str) -> bool:
    return _ext(path) in EXCEL_EXTENSIONS


def _is_jsonl(path: str) -> bool:
    return _ext(path) in JSONL_EXTENSIONS


def _detect_type_by_path(path: str) -> DataType:
    if not os.path.exists(path):
        return DataType.UNKNOWN
    if os.path.isdir(path):
        return DataType.DIRECTORY
    ext = _ext(path)
    if ext in IMAGE_EXTENSIONS:
        return DataType.IMAGE
    if ext in CSV_EXTENSIONS:
        return DataType.CSV
    if ext in EXCEL_EXTENSIONS:
        return DataType.EXCEL
    if ext in JSONL_EXTENSIONS:
        return DataType.JSONL
    if ext in TEXT_EXTENSIONS:
        return DataType.TEXT
    # 内容嗅探
    try:
        with open(path, "r", encoding="utf-8") as f:
            head = f.read(1024)
        if "," in head:
            lines = head.splitlines()
            if len(lines) > 1:
                nc0 = len(lines[0].split(","))
                nc1 = len(lines[1].split(",")) if lines[1] else 0
                if nc0 > 2 and nc0 == nc1:
                    return DataType.CSV
        return DataType.TEXT
    except UnicodeDecodeError:
        return DataType.IMAGE
    except Exception:
        return DataType.UNKNOWN


def _read_csv_safe(path: str, n_rows: int = 5):
    """读取 CSV 文件，自动处理编码。返回 (columns, rows, total_rows)。"""
    for enc in ["utf-8", "utf-8-sig", "gbk", "gb2312", "latin-1"]:
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                total = sum(1 for _ in f) - 1
            with open(path, "r", encoding=enc, newline="") as f:
                reader = csv.reader(f)
                cols = next(reader, None)
                if not cols:
                    continue
                rows = [r for _, r in zip(range(n_rows), reader)]
            return cols, rows, max(total, 0)
        except UnicodeDecodeError:
            continue
        except Exception:
            break
    return None, [], 0


def _analyze_columns(columns: List[str], sample_rows: List[List[str]]) -> Dict[str, Any]:
    """基于关键词和内容对列进行分类分析。"""
    kw = MEDICAL_COL_KEYWORDS

    std_cols, val_cols, unit_cols, text_cols = [], [], [], []

    for i, col in enumerate(columns):
        col_lower = col.lower()

        # 长文本列检测（按列名关键词 + 内容长度）
        is_text_kw = any(k in col_lower for k in kw["extract_text"])
        col_values = [r[i] for r in sample_rows if i < len(r) and r[i]]
        avg_len = sum(len(v) for v in col_values) / max(len(col_values), 1)
        if is_text_kw and avg_len > 80:
            text_cols.append(col)
            continue

        # 术语标准化候选
        if any(k in col_lower for k in kw["standardize"]):
            std_cols.append(col)
            continue

        # 数值量纲候选
        if any(k in col_lower for k in kw["unit_normalize"]):
            numeric_count = 0
            for v in col_values:
                try:
                    float(v)
                    numeric_count += 1
                except (ValueError, TypeError):
                    pass
            if col_values and numeric_count / len(col_values) > 0.5:
                val_cols.append(col)
            continue

        # 单位列
        if any(k in col_lower for k in kw["unit_col"]):
            unit_cols.append(col)

    # 自动匹配 value->unit
    unit_mapping: Dict[str, str] = {}
    for vc in val_cols:
        base = re.sub(r"_value$|_result$|_amount$", "", vc.lower())
        for uc in unit_cols:
            if base in uc.lower() or uc.lower().replace("_text", "").replace("_unit", "") in base:
                unit_mapping[vc] = uc
                break

    return {
        "standardization_candidates": std_cols,
        "value_columns": val_cols,
        "unit_columns": unit_cols,
        "unit_mapping": unit_mapping,
        "extraction_candidates": text_cols,
    }


# ---------------------------------------------------------------------------
# 公开工具函数（每个函数都设计为可直接注册到 Toolkit 的原子操作）
# ---------------------------------------------------------------------------

def detect_input_type(input_path: str) -> Dict[str, Any]:
    """
    检测输入路径的数据类型及基本元信息。

    Args:
        input_path: 文件路径或目录路径

    Returns:
        {
          "data_type": "image|text|csv|directory|unknown",
          "exists": bool,
          "is_dir": bool,
          "extension": str,
          "file_name": str,
          "file_size_kb": float,
        }
    """
    result: Dict[str, Any] = {
        "data_type": DataType.UNKNOWN.value,
        "exists": os.path.exists(input_path),
        "is_dir": False,
        "extension": _ext(input_path),
        "file_name": os.path.basename(input_path),
        "file_size_kb": 0.0,
    }
    if not result["exists"]:
        return result
    if os.path.isdir(input_path):
        result["is_dir"] = True
        result["data_type"] = DataType.DIRECTORY.value
        return result
    result["file_size_kb"] = round(os.path.getsize(input_path) / 1024, 2)
    result["data_type"] = _detect_type_by_path(input_path).value
    return result


def read_csv_sample(
    file_path: str,
    sample_rows: int = 5,
) -> Dict[str, Any]:
    """
    读取 CSV 文件的列名和前 N 行样本，同时返回基于规则的列分析结果。

    Args:
        file_path: CSV 文件路径
        sample_rows: 采样行数（默认 5）

    Returns:
        {
          "success": bool,
          "columns": [列名列表],
          "sample_data": [[值,...], ...],   # 列表的列表（按列索引对应）
          "sample_dict": [{列名:值,...},...], # 字典形式，便于阅读
          "total_rows": int,
          "column_analysis": {              # 自动列分析
            "standardization_candidates": [列名],
            "value_columns": [列名],
            "unit_columns": [列名],
            "unit_mapping": {值列: 单位列},
            "extraction_candidates": [列名],
          },
          "error": str | None,
        }
    """
    cols, rows, total = _read_csv_safe(file_path, sample_rows)
    if cols is None:
        return {"success": False, "columns": [], "sample_data": [], "sample_dict": [],
                "total_rows": 0, "column_analysis": {}, "error": "无法读取 CSV 文件"}

    sample_dict = [dict(zip(cols, r)) for r in rows]
    analysis = _analyze_columns(cols, rows)

    return {
        "success": True,
        "columns": cols,
        "sample_data": rows,
        "sample_dict": sample_dict,
        "total_rows": total,
        "column_analysis": analysis,
        "error": None,
    }


def get_csv_full_info(file_path: str) -> str:
    """
    获取 CSV 文件的完整可读分析报告（字符串形式），供 LLM 推理参考。

    Args:
        file_path: CSV 文件路径

    Returns:
        包含文件信息、列名、样本数据和候选处理列的文本报告
    """
    info = read_csv_sample(file_path, sample_rows=3)
    if not info["success"]:
        return f"[错误] 无法读取 CSV: {info.get('error', '未知错误')}"

    cols = info["columns"]
    rows = info["sample_dict"]
    analysis = info["column_analysis"]
    total = info["total_rows"]

    lines = [
        f"=== CSV 文件分析报告 ===",
        f"文件: {os.path.basename(file_path)}",
        f"总行数: {total}  |  列数: {len(cols)}",
        "",
        f"列名列表:",
        "  " + ", ".join(cols),
        "",
        f"样本数据（前 {len(rows)} 行）:",
    ]
    for i, row in enumerate(rows, 1):
        row_str = "  |  ".join(f"{k}: {v!r}" for k, v in list(row.items())[:8])
        lines.append(f"  [{i}] {row_str}")

    lines += [
        "",
        "--- 自动列分析结果 ---",
        f"  术语标准化候选: {analysis.get('standardization_candidates', [])}",
        f"  数值量纲候选:   {analysis.get('value_columns', [])}",
        f"  单位列:         {analysis.get('unit_columns', [])}",
        f"  值→单位映射:    {analysis.get('unit_mapping', {})}",
        f"  长文本抽取候选: {analysis.get('extraction_candidates', [])}",
    ]
    return "\n".join(lines)


def scan_directory(folder_path: str) -> Dict[str, Any]:
    """
    扫描目录结构，识别患者 ID 分组和文件类型分布。

    目录规范：根目录 / 患者ID / 类别子目录 / 文件
    患者 ID 规则：4–8 位纯数字的子目录名。

    Args:
        folder_path: 目录路径

    Returns:
        {
          "success": bool,
          "folder": str,
          "patients": {
            患者ID: {
              "file_count": int,
              "categories": [类别名],
              "files": [{"path": ..., "category": ..., "file_type": ...}]
            }
          },
          "statistics": {
            "total_patients": int,
            "total_files": int,
            "image_files": int,
            "text_files": int,
          },
          "has_table": bool,       # 是否存在 Excel 表格
          "table_path": str|None,
          "summary": str,          # 人类可读摘要
          "error": str|None,
        }
    """
    if not os.path.exists(folder_path):
        return {"success": False, "error": f"路径不存在: {folder_path}"}
    if not os.path.isdir(folder_path):
        return {"success": False, "error": f"不是目录: {folder_path}"}

    skip = get_skip_folders()
    # "table" 子目录存放参考表格，不作为待处理的数据文件扫描
    file_scan_skip = skip | {"table"}
    patient_re = re.compile(r"^\d{4,8}$")
    patients: Dict[str, Any] = {}
    stats = {"total_patients": 0, "total_files": 0, "image_files": 0, "text_files": 0}
    table_path: Optional[str] = None
    _table_exts = (".xlsx", ".xls", ".csv")

    for entry in sorted(os.listdir(folder_path)):
        full = os.path.join(folder_path, entry)
        # 检测顶层 Excel/CSV 表格
        if os.path.isfile(full) and entry.endswith(_table_exts):
            if table_path is None:
                table_path = full
            continue
        # 检测 table 子目录（递归查找第一个 xlsx/csv）
        if os.path.isdir(full) and entry.lower() == "table" and table_path is None:
            for dirpath, _, fnames in os.walk(full):
                for fname in sorted(fnames):
                    if fname.endswith(_table_exts):
                        table_path = os.path.join(dirpath, fname)
                        break
                if table_path:
                    break
            continue
        # 只处理患者 ID 目录
        if not (os.path.isdir(full) and patient_re.match(entry)):
            continue

        pid = entry
        patients[pid] = {"file_count": 0, "categories": [], "files": [], "table_path": None}

        for cat in sorted(os.listdir(full)):
            cat_path = os.path.join(full, cat)
            if not os.path.isdir(cat_path):
                continue
            # 检测患者级 table 子目录
            if cat.lower() == "table":
                # 根目录 CSV（作为主表格路径）
                for fname in sorted(os.listdir(cat_path)):
                    fpath_full = os.path.join(cat_path, fname)
                    if os.path.isfile(fpath_full) and fname.endswith(_table_exts):
                        patients[pid]["table_path"] = fpath_full
                        if table_path is None:
                            table_path = fpath_full
                        break
                # 扫描子目录（notes/, structured/ 等）下的 CSV
                table_sub_files: List[Dict[str, str]] = []
                for subentry in sorted(os.listdir(cat_path)):
                    sub_path = os.path.join(cat_path, subentry)
                    if os.path.isdir(sub_path):
                        for fname in sorted(os.listdir(sub_path)):
                            if fname.lower().endswith(_table_exts):
                                table_sub_files.append({
                                    "path": os.path.join(sub_path, fname),
                                    "subdir": subentry,
                                    "filename": fname,
                                })
                patients[pid]["table_sub_files"] = table_sub_files
                continue
            if cat in file_scan_skip:
                continue
            patients[pid]["categories"].append(cat)

            for dirpath, _, fnames in os.walk(cat_path):
                for fname in sorted(fnames):
                    fpath = os.path.join(dirpath, fname)
                    ftype = _detect_type_by_path(fpath).value
                    patients[pid]["files"].append({
                        "path": fpath,
                        "category": cat,
                        "file_type": ftype,
                        "filename": fname,
                    })
                    patients[pid]["file_count"] += 1
                    stats["total_files"] += 1
                    if ftype == "image":
                        stats["image_files"] += 1
                    elif ftype == "text":
                        stats["text_files"] += 1

    stats["total_patients"] = len(patients)

    # 收集根目录下的非表格数据文件（JSONL、CSV、文本等）
    flat_files: List[Dict[str, Any]] = []
    for entry in sorted(os.listdir(folder_path)):
        full = os.path.join(folder_path, entry)
        if not os.path.isfile(full):
            continue
        if entry.endswith(_table_exts) and full == table_path:
            continue  # 已作为表格，跳过
        ftype = _detect_type_by_path(full).value
        if ftype not in ("unknown",):
            flat_files.append({"path": full, "filename": entry, "file_type": ftype})

    # 人类可读摘要
    summary_lines = [
        f"目录: {folder_path}",
        f"患者数: {stats['total_patients']}  |  文件总数: {stats['total_files']}",
        f"  图片: {stats['image_files']}  文本: {stats['text_files']}",
    ]
    if table_path:
        summary_lines.append(f"  Excel 表格: {os.path.basename(table_path)}")
    if flat_files:
        summary_lines.append(f"  根目录数据文件: {[f['filename'] for f in flat_files]}")
    for pid, pdata in list(patients.items())[:5]:
        sub_info = ""
        if pdata.get("table_sub_files"):
            sub_dirs: Dict[str, List[str]] = {}
            for sf in pdata["table_sub_files"]:
                sub_dirs.setdefault(sf["subdir"], []).append(sf["filename"])
            sub_info = ", table子目录: " + "; ".join(
                f"{d}({', '.join(fs)})" for d, fs in sub_dirs.items()
            )
        table_info = f", 表格: {os.path.basename(pdata['table_path'])}" if pdata.get("table_path") else ""
        summary_lines.append(
            f"  患者 {pid}: {pdata['file_count']} 个文件, "
            f"类别: {pdata['categories']}"
            + table_info + sub_info
        )
    if len(patients) > 5:
        summary_lines.append(f"  ... 还有 {len(patients) - 5} 个患者")

    return {
        "success": True,
        "folder": folder_path,
        "patients": patients,
        "flat_files": flat_files,
        "statistics": stats,
        "has_table": table_path is not None,
        "table_path": table_path,
        "summary": "\n".join(summary_lines),
        "error": None,
    }


def list_directory(folder_path: str) -> Dict[str, Any]:
    """
    列出目录下的直接子文件和子目录（非递归）。

    Args:
        folder_path: 目录路径

    Returns:
        {"files": [文件名], "subdirs": [子目录名], "total": int, "error": str|None}
    """
    if not os.path.exists(folder_path):
        return {"files": [], "subdirs": [], "total": 0, "error": f"路径不存在: {folder_path}"}
    if not os.path.isdir(folder_path):
        return {"files": [], "subdirs": [], "total": 0, "error": f"不是目录: {folder_path}"}

    files, subdirs = [], []
    for entry in sorted(os.listdir(folder_path)):
        full = os.path.join(folder_path, entry)
        (subdirs if os.path.isdir(full) else files).append(entry)

    return {"files": files, "subdirs": subdirs, "total": len(files) + len(subdirs), "error": None}


def read_text_sample(file_path: str, max_chars: int = 2000) -> Dict[str, Any]:
    """
    读取文本文件的前 N 个字符，供 LLM 判断文本内容类型。

    Args:
        file_path: 文本文件路径
        max_chars: 最多读取字符数（默认 2000）

    Returns:
        {"success": bool, "content": str, "file_size_kb": float, "encoding": str, "error": str|None}
    """
    if not os.path.exists(file_path):
        return {"success": False, "content": "", "file_size_kb": 0.0,
                "encoding": "", "error": f"文件不存在: {file_path}"}

    size_kb = round(os.path.getsize(file_path) / 1024, 2)
    for enc in ["utf-8", "gbk", "latin-1"]:
        try:
            with open(file_path, "r", encoding=enc) as f:
                content = f.read(max_chars)
            return {"success": True, "content": content, "file_size_kb": size_kb,
                    "encoding": enc, "error": None}
        except UnicodeDecodeError:
            continue

    return {"success": False, "content": "", "file_size_kb": size_kb,
            "encoding": "", "error": "无法解析文件编码"}


def read_excel_sample(
    file_path: str,
    sample_rows: int = 5,
    sheet_index: int = 0,
) -> Dict[str, Any]:
    """
    读取 Excel 文件（.xlsx/.xls）的列名和前 N 行样本，同时返回自动列分析结果。

    Args:
        file_path: Excel 文件路径
        sample_rows: 采样行数（默认 5）
        sheet_index: 工作表索引（默认 0，即第一张表）

    Returns:
        {
          "success": bool,
          "columns": [列名列表],
          "sample_dict": [{列名:值,...},...],
          "total_rows": int,
          "sheet_name": str,
          "column_analysis": {...},
          "error": str | None,
        }
    """
    if not os.path.exists(file_path):
        return {"success": False, "columns": [], "sample_dict": [], "total_rows": 0,
                "sheet_name": "", "column_analysis": {}, "error": f"文件不存在: {file_path}"}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet_name = wb.sheetnames[sheet_index] if sheet_index < len(wb.sheetnames) else wb.sheetnames[0]
        ws = wb[sheet_name]

        rows_iter = ws.iter_rows(values_only=True)
        header_raw = next(rows_iter, None)
        if not header_raw:
            return {"success": False, "columns": [], "sample_dict": [], "total_rows": 0,
                    "sheet_name": sheet_name, "column_analysis": {}, "error": "Excel 文件无表头"}

        columns = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(header_raw)]
        sample_list: List[List[str]] = []
        total = 0
        for row in rows_iter:
            total += 1
            if len(sample_list) < sample_rows:
                sample_list.append([str(v) if v is not None else "" for v in row])

        sample_dict = [dict(zip(columns, r)) for r in sample_list]
        analysis = _analyze_columns(columns, sample_list)

        return {
            "success": True,
            "columns": columns,
            "sample_dict": sample_dict,
            "total_rows": total,
            "sheet_name": sheet_name,
            "column_analysis": analysis,
            "error": None,
        }
    except Exception as e:
        return {"success": False, "columns": [], "sample_dict": [], "total_rows": 0,
                "sheet_name": "", "column_analysis": {}, "error": str(e)}


def read_jsonl_sample(file_path: str, sample_rows: int = 3) -> Dict[str, Any]:
    """
    读取 JSONL 文件的前 N 条记录，返回顶层字段名和样本摘要。

    Args:
        file_path: .jsonl 文件路径
        sample_rows: 采样条数（默认 3）

    Returns:
        {
          "success": bool,
          "total_rows": int,
          "top_keys": [顶层字段名],
          "sample": [dict, ...],   # 前 N 条记录（仅保留顶层字段摘要）
          "error": str | None,
        }
    """
    if not os.path.exists(file_path):
        return {"success": False, "total_rows": 0, "top_keys": [], "sample": [],
                "error": f"文件不存在: {file_path}"}
    try:
        records = []
        total = 0
        with open(file_path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    obj = json.loads(line)
                    total += 1
                    if len(records) < sample_rows:
                        records.append(obj)
                except json.JSONDecodeError:
                    continue

        top_keys = list(records[0].keys()) if records else []

        # 生成摘要：对每条记录只保留顶层字段的类型/长度信息
        sample_summary = []
        for rec in records:
            summary = {}
            for k, v in rec.items():
                if isinstance(v, dict):
                    summary[k] = f"{{dict, keys={list(v.keys())[:5]}}}"
                elif isinstance(v, list):
                    summary[k] = f"[list, len={len(v)}]"
                elif isinstance(v, str) and len(v) > 80:
                    summary[k] = f"[长文本 {len(v)}字] {v[:60]}..."
                else:
                    summary[k] = v
            sample_summary.append(summary)

        return {
            "success": True,
            "total_rows": total,
            "top_keys": top_keys,
            "sample": sample_summary,
            "error": None,
        }
    except Exception as e:
        return {"success": False, "total_rows": 0, "top_keys": [], "sample": [], "error": str(e)}
